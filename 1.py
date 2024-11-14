import os
import re
import tkinter as tk
from tkinter import messagebox, ttk, filedialog
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment
from difflib import SequenceMatcher
import pandas as pd

def extract_data_from_excel_dynamic(excel_path):
    df = pd.read_excel(excel_path)
    df['试题题目'] = df['试题题目'].apply(lambda x: clean_text(x))

    classifications_columns_indices = [14, 15, 16]  # 依然使用 O, P, Q 列的索引

    data = {
        "single_choice": [],
        "multiple_choice": [],
        "true_false": [],
        "single_classifications": [],
        "multiple_classifications": [],
        "true_false_classifications": [],
        "single_choice_options": [],
        "multiple_choice_options": [],
        "true_false_options": [],
        "single_choice_correct_answers": [],
        "multiple_choice_correct_answers": [],
        "true_false_correct_answers": [],
        "single_choice_map": {},
        "multiple_choice_map": {},
        "true_false_map": {}
    }

    for index, row in df.iterrows():
        question_type = row['题型']
        question_content = row['试题题目']
        options = [row['A'], row['B'], row['C'], row['D'], row['E'], row['F'], row['G'], row['H']]
        correct_answer = row['答案']

        classifications = [row.iloc[classifications_columns_indices[0]],
                           row.iloc[classifications_columns_indices[1]],
                           row.iloc[classifications_columns_indices[2]]]  # 使用方括号 []

        if question_type == "单选题":
            data["single_choice"].append(question_content)
            data["single_classifications"].append(classifications)
            data["single_choice_options"].append(options)
            data["single_choice_correct_answers"].append(correct_answer)
            data["single_choice_map"][question_content] = len(data["single_choice"]) - 1

        elif question_type == "多选题":
            data["multiple_choice"].append(question_content)
            data["multiple_classifications"].append(classifications)
            data["multiple_choice_options"].append(options)
            data["multiple_choice_correct_answers"].append(correct_answer)
            data["multiple_choice_map"][question_content] = len(data["multiple_choice"]) - 1

        elif question_type == "判断题":
            data["true_false"].append(question_content)
            data["true_false_classifications"].append(classifications)
            data["true_false_options"].append(["正确", "错误"])
            data["true_false_correct_answers"].append(correct_answer)
            data["true_false_map"][question_content] = len(data["true_false"]) - 1

    return data


def clean_text(text):
    """
    清理文本内容，去除空白字符、换行符、多余空格以及不可见字符。
    """
    if isinstance(text, str):
        # 替换各种换行符为空格
        text = text.replace('\n', ' ').replace('\r', ' ')

        # 移除不可见字符（包括不可打印字符和零宽度空格）
        text = ''.join(ch for ch in text if ch.isprintable())

        # 移除多余的空格（将多个连续的空格替换为单个空格）
        text = re.sub(r'\s+', ' ', text).strip()

    return text


def get_best_match(word, possibilities, threshold=0.5):
    """
    先进行正序匹配，如果没有匹配到合适的项，再进行倒序匹配。
    """
    best_match = None
    best_score = threshold

    # 正序匹配
    for possibility in possibilities:
        score = SequenceMatcher(None, word, possibility).ratio()
        if score > best_score:
            best_score = score
            best_match = possibility

    # 如果正序匹配没有找到足够好的匹配项，尝试倒序匹配
    if best_match is None or best_score < threshold:
        reversed_word = word[::-1]
        for possibility in possibilities:
            reversed_possibility = possibility[::-1]
            reversed_score = SequenceMatcher(None, reversed_word, reversed_possibility).ratio()
            if reversed_score > best_score:
                best_score = reversed_score
                best_match = possibility

    return best_match, best_score

def extract_data_from_docx(file_path, excel_data):
    doc = Document(file_path)
    data = {
        "name": "",
        "total_score": 0.0,
        "single_choice": [None] * len(excel_data['single_choice']),
        "multiple_choice": [None] * len(excel_data['multiple_choice']),
        "true_false": [None] * len(excel_data['true_false'])
    }

    question_pattern = re.compile(r"^\d+\.\s*(.*)$")
    current_question_content = None

    for para in doc.paragraphs:
        text = para.text.strip()
        name_match = re.search(r"考生名称：([^\n]+)", text)

        if name_match:
            data["name"] = name_match.group(1).strip()

        question_match = question_pattern.match(text)
        if question_match:
            current_question_content = question_match.group(1).strip()
            cleaned_question_content = clean_text(current_question_content)

        if current_question_content and "该题得分是:" in text:
            score = float(text.split("该题得分是:")[1].split()[0].replace("分", "").strip())
            data["total_score"] += score

            found_match = False
            if cleaned_question_content in excel_data["single_choice_map"]:
                index = excel_data["single_choice_map"][cleaned_question_content]
                data["single_choice"][index] = score
                found_match = True
            elif cleaned_question_content in excel_data["multiple_choice_map"]:
                index = excel_data["multiple_choice_map"][cleaned_question_content]
                data["multiple_choice"][index] = score
                found_match = True
            elif cleaned_question_content in excel_data["true_false_map"]:
                index = excel_data["true_false_map"][cleaned_question_content]
                data["true_false"][index] = score
                found_match = True

            if not found_match:
                all_questions = list(excel_data['single_choice_map'].keys()) + \
                                list(excel_data['multiple_choice_map'].keys()) + \
                                list(excel_data['true_false_map'].keys())

                best_match, best_score = get_best_match(cleaned_question_content, all_questions, threshold=0.5)
                if best_match:
                    if best_match in excel_data["single_choice_map"]:
                        index = excel_data["single_choice_map"][best_match]
                        data["single_choice"][index] = score
                    elif best_match in excel_data["multiple_choice_map"]:
                        index = excel_data["multiple_choice_map"][best_match]
                        data["multiple_choice"][index] = score
                    elif best_match in excel_data["true_false_map"]:
                        index = excel_data["true_false_map"][best_match]
                        data["true_false"][index] = score

            current_question_content = None

    # Round the total score to one decimal place
    data["total_score"] = round(data["total_score"], 1)

    return data

def save_to_excel(data_list, excel_data, save_path):
    if os.path.exists(save_path):
        replace = messagebox.askyesno("文件已存在", f"文件 '{save_path}' 已存在。是否替换？")
        if not replace:
            messagebox.showinfo("操作取消", "保存操作已取消。")
            return

    wb = Workbook()
    ws = wb.active

    # 设置表头
    headers = [
        "题目类型", "题目内容",
        "答案选项A", "答案选项B", "答案选项C", "答案选项D", "答案选项E", "答案选项F", "答案选项G", "答案选项H",
        "正确答案",
        "一级分类", "二级分类", "三级分类"
    ]

    for data in data_list:
        headers.append(f"{data['name']}/{data['total_score']}")

    headers.append("得分率")
    ws.append(headers)

    # 添加单选题数据
    for i, question in enumerate(excel_data['single_choice']):
        row = [
            "单选题",
            question,
            excel_data['single_choice_options'][i][0],  # 答案选项A
            excel_data['single_choice_options'][i][1],  # 答案选项B
            excel_data['single_choice_options'][i][2],  # 答案选项C
            excel_data['single_choice_options'][i][3],  # 答案选项D
            excel_data['single_choice_options'][i][4],  # 答案选项E
            excel_data['single_choice_options'][i][5],  # 答案选项F
            excel_data['single_choice_options'][i][6],  # 答案选项G
            excel_data['single_choice_options'][i][7],  # 答案选项H
            excel_data['single_choice_correct_answers'][i],  # 正确答案
            excel_data['single_classifications'][i][0],  # 一级分类
            excel_data['single_classifications'][i][1],  # 二级分类
            excel_data['single_classifications'][i][2],  # 三级分类
        ]
        scores = []
        for data in data_list:
            score = data['single_choice'][i]
            row.append(score)
            scores.append(score)

        non_zero_scores = [score for score in scores if score is not None and score > 0]
        score_rate = len(non_zero_scores) / len(scores) if scores else 0
        score_rate_percentage = "{:.0%}".format(score_rate)
        row.append(score_rate_percentage)

        ws.append(row)

    # 添加多选题数据
    for i, question in enumerate(excel_data['multiple_choice']):
        row = [
            "多选题",
            question,
            excel_data['multiple_choice_options'][i][0],  # 答案选项A
            excel_data['multiple_choice_options'][i][1],  # 答案选项B
            excel_data['multiple_choice_options'][i][2],  # 答案选项C
            excel_data['multiple_choice_options'][i][3],  # 答案选项D
            excel_data['multiple_choice_options'][i][4],  # 答案选项E
            excel_data['multiple_choice_options'][i][5],  # 答案选项F
            excel_data['multiple_choice_options'][i][6],  # 答案选项G
            excel_data['multiple_choice_options'][i][7],  # 答案选项H
            excel_data['multiple_choice_correct_answers'][i],  # 正确答案
            excel_data['multiple_classifications'][i][0],  # 一级分类
            excel_data['multiple_classifications'][i][1],  # 二级分类
            excel_data['multiple_classifications'][i][2],  # 三级分类
        ]
        scores = []
        for data in data_list:
            score = data['multiple_choice'][i]
            row.append(score)
            scores.append(score)

        non_zero_scores = [score for score in scores if score is not None and score > 0]
        score_rate = len(non_zero_scores) / len(scores) if scores else 0
        score_rate_percentage = "{:.0%}".format(score_rate)
        row.append(score_rate_percentage)

        ws.append(row)

    # 添加判断题数据
    for i, question in enumerate(excel_data['true_false']):
        row = [
            "判断题",
            question,
            "正确",  # 答案选项A
            "错误",  # 答案选项B
            "", "", "", "", "", "",  # 判断题没有其他选项
            excel_data['true_false_correct_answers'][i],  # 正确答案
            excel_data['true_false_classifications'][i][0],  # 一级分类
            excel_data['true_false_classifications'][i][1],  # 二级分类
            excel_data['true_false_classifications'][i][2],  # 三级分类
        ]
        scores = []
        for data in data_list:
            score = data['true_false'][i]
            row.append(score)
            scores.append(score)

        non_zero_scores = [score for score in scores if score is not None and score > 0]
        score_rate = len(non_zero_scores) / len(scores) if scores else 0
        score_rate_percentage = "{:.0%}".format(score_rate)
        row.append(score_rate_percentage)

        ws.append(row)

    alignment_center = Alignment(horizontal='center', vertical='center')
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = alignment_center

    try:
        wb.save(save_path)
        messagebox.showinfo("完成", f"成绩表格已保存为 '{save_path}'！")
    except PermissionError:
        messagebox.showerror("错误", f"无法保存文件 '{save_path}'。请确保文件未被其他程序打开，然后重试。")

def process_files(files, excel_data):
    data_list = []
    for file in files:
        data = extract_data_from_docx(file, excel_data)
        data_list.append(data)

    save_to_excel(data_list, excel_data, "提取后的试卷分析.xlsx")

def select_files():
    """
    使用文件对话框让用户选择Word文件，并处理选择的文件。
    """
    file_paths = filedialog.askopenfilenames(filetypes=[("Word files", "*.docx")])
    if file_paths:
        # 先抓取Excel数据
        excel_data = extract_data_from_excel_dynamic("数据源表格.xlsx")
        process_files(file_paths, excel_data)

# 创建GUI窗口
root = tk.Tk()
root.title("Word to Excel Converter")  # 设置窗口标题
root.geometry('800x600')  # 调整窗口大小为800x600

# 创建一个大区域用于文件选择
frame = ttk.Frame(root)
frame.pack(expand=True, fill='both', padx=20, pady=20)

label = ttk.Label(frame, text="请点击下方按钮选择Word文件", background="lightgray", anchor="center")
label.pack(expand=True, fill='both')

# 添加按钮用于重新选择文件
select_button = ttk.Button(root, text="选择文件", command=select_files)
select_button.pack(pady=10)

root.mainloop()
